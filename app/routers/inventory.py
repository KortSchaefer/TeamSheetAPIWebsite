import csv
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from html import escape
from io import BytesIO, StringIO

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import HTMLResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload, selectinload

from app import schemas
from app.core.security import get_current_manager_or_admin, get_current_user
from app.database import get_db
from app.models import (
    InventoryBalance,
    InventoryCount,
    InventoryCountLine,
    InventoryCountStatus,
    InventoryCountTemplate,
    InventoryCountTemplateLine,
    InventoryVoiceEntry,
    InventoryVoiceReviewStatus,
    InventoryVoiceSessionCount,
    InventoryVoiceSessionStatus,
    InventoryItem,
    InventoryLocation,
    InventoryWeekdayTarget,
    InventoryReceiving,
    InventoryReceivingLine,
    Ingredient,
    PurchaseOrder,
    PurchaseOrderLine,
    PurchaseOrderStatus,
    StockMovement,
    User,
    Vendor,
    VendorItem,
)
from app.services.count_sheets import (
    apply_line_changes,
    populate_count_sheet,
    serialize_count_sheet,
    serialize_template,
)
from app.services.purchase_order_import import (
    PurchaseOrderCSVError,
    preview_purchase_order_csv,
)
from app.services.inventory_planning import (
    inventory_settings_rows,
    purchase_order_planner_rows,
)

router = APIRouter(prefix="/inventory", tags=["inventory"])


def get_location(db: Session, location_id: int) -> InventoryLocation:
    location = db.query(InventoryLocation).filter(InventoryLocation.id == location_id).first()
    if not location:
        raise HTTPException(status_code=404, detail="Inventory location not found")
    return location


def get_item(db: Session, item_id: int) -> InventoryItem:
    item = db.query(InventoryItem).filter(InventoryItem.id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Inventory item not found")
    return item


def get_or_create_balance(db: Session, item_id: int, location_id: int) -> InventoryBalance:
    balance = (
        db.query(InventoryBalance)
        .filter(InventoryBalance.inventory_item_id == item_id, InventoryBalance.location_id == location_id)
        .first()
    )
    if balance:
        return balance
    balance = InventoryBalance(inventory_item_id=item_id, location_id=location_id)
    db.add(balance)
    db.flush()
    return balance


def movement_response(movement: StockMovement) -> schemas.InventoryMovementRead:
    return schemas.InventoryMovementRead(
        id=movement.id,
        inventory_item_id=movement.inventory_item_id,
        location_id=movement.location_id,
        quantity_change=Decimal(str(movement.quantity_change)),
        reason=movement.reason,
        notes=movement.notes,
        source_event_key=movement.source_event_key,
        lot_number=movement.lot_number,
        expiration_date=movement.expiration_date,
        created_by_user_id=movement.created_by_user_id,
        created_at=movement.created_at,
        updated_at=movement.updated_at,
    )


def purchase_order_response(order: PurchaseOrder) -> dict:
    return {
        "id": order.id,
        "vendor_id": order.vendor_id,
        "status": order.status,
        "expected_date": order.expected_date,
        "notes": order.notes,
        "external_reference": order.external_reference,
        "imported_filename": order.imported_filename,
        "created_by_user_id": order.created_by_user_id,
        "lines": [
            {
                "id": line.id,
                "inventory_item_id": line.inventory_item_id,
                "item_name": line.item.name if line.item else None,
                "location_id": line.location_id,
                "location_name": line.location.name if line.location else None,
                "ordered_quantity": line.ordered_quantity,
                "received_quantity": line.received_quantity,
                "remaining_quantity": max(
                    Decimal(line.ordered_quantity or 0)
                    - Decimal(line.received_quantity or 0),
                    Decimal("0"),
                ),
                "unit_price_cents": line.unit_price_cents,
                "purchase_unit": line.purchase_unit or line.item.purchase_unit or line.item.base_unit,
                "quantity_per_purchase_unit": line.quantity_per_purchase_unit,
            }
            for line in order.lines
        ],
    }


def ensure_voice_count_ready(db: Session, count_id: int) -> None:
    link = (
        db.query(InventoryVoiceSessionCount)
        .filter(InventoryVoiceSessionCount.inventory_count_id == count_id)
        .first()
    )
    if not link:
        return
    blocking = (
        db.query(InventoryVoiceEntry)
        .filter(
            InventoryVoiceEntry.session_id == link.session_id,
            InventoryVoiceEntry.review_status
            == InventoryVoiceReviewStatus.NEEDS_REVIEW,
        )
        .count()
    )
    if blocking or link.session.status != InventoryVoiceSessionStatus.FINISHED:
        raise HTTPException(
            status_code=409,
            detail=(
                f"Voice inventory session #{link.session_id} has "
                f"{blocking} unresolved entr{'y' if blocking == 1 else 'ies'}"
            ),
        )


def ensure_count_sheet_complete(count: InventoryCount) -> None:
    uncounted = sum(1 for line in count.lines if not line.is_counted)
    exceptions = sum(
        1 for line in count.lines if line.review_status == "NEEDS_REVIEW"
    )
    if uncounted or exceptions:
        raise HTTPException(
            status_code=409,
            detail=(
                f"Count sheet has {uncounted} uncounted row(s) and "
                f"{exceptions} unresolved exception(s)"
            ),
        )


@router.get("/locations", response_model=list[schemas.InventoryLocationRead])
def list_locations(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return db.query(InventoryLocation).filter(InventoryLocation.active.is_(True)).order_by(InventoryLocation.name).all()


@router.post("/locations", response_model=schemas.InventoryLocationRead, status_code=status.HTTP_201_CREATED)
def create_location(
    payload: schemas.InventoryLocationCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_manager_or_admin),
):
    location = InventoryLocation(**payload.model_dump())
    db.add(location)
    db.commit()
    db.refresh(location)
    return location


@router.get("/items", response_model=list[schemas.InventoryItemRead])
def list_items(
    active: bool | None = Query(default=None),
    search: str | None = Query(default=None),
    category: str | None = Query(default=None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(InventoryItem)
    if active is not None:
        query = query.filter(InventoryItem.active == active)
    if search:
        query = query.filter(InventoryItem.name.ilike(f"%{search}%"))
    if category:
        query = query.filter(InventoryItem.category == category)
    return query.order_by(InventoryItem.category, InventoryItem.name).all()


@router.post("/items", response_model=schemas.InventoryItemRead, status_code=status.HTTP_201_CREATED)
def create_item(
    payload: schemas.InventoryItemCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_manager_or_admin),
):
    if payload.default_location_id:
        get_location(db, payload.default_location_id)
    item = InventoryItem(**payload.model_dump())
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


@router.post(
    "/items/from-catalog",
    response_model=schemas.InventoryItemRead,
    status_code=status.HTTP_201_CREATED,
)
def activate_catalog_item(
    payload: schemas.InventoryCatalogActivationCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_manager_or_admin),
):
    ingredient = (
        db.query(Ingredient)
        .filter(Ingredient.external_id == payload.catalog_id, Ingredient.active.is_(True))
        .first()
    )
    if ingredient is None:
        raise HTTPException(status_code=404, detail="Catalog ingredient not found")
    if payload.default_location_id:
        get_location(db, payload.default_location_id)

    item = (
        db.query(InventoryItem)
        .filter(InventoryItem.ingredient_id == ingredient.id)
        .first()
    )
    values = payload.model_dump(exclude={"catalog_id"})
    if item is None:
        item = InventoryItem(
            ingredient_id=ingredient.id,
            name=ingredient.name,
            category=ingredient.category,
            active=True,
            **values,
        )
        db.add(item)
    else:
        item.active = True
        for key, value in values.items():
            setattr(item, key, value)
    db.commit()
    db.refresh(item)
    return item


@router.put("/items/{item_id}", response_model=schemas.InventoryItemRead)
def update_item(
    item_id: int,
    payload: schemas.InventoryItemCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_manager_or_admin),
):
    item = get_item(db, item_id)
    if payload.default_location_id:
        get_location(db, payload.default_location_id)
    for key, value in payload.model_dump().items():
        setattr(item, key, value)
    db.commit()
    db.refresh(item)
    return item


@router.get("/items/{item_id}/balances", response_model=list[schemas.InventoryBalanceRead])
def item_balances(item_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    get_item(db, item_id)
    return db.query(InventoryBalance).filter(InventoryBalance.inventory_item_id == item_id).all()


@router.put("/items/{item_id}/balances", response_model=schemas.InventoryBalanceRead)
def upsert_balance(
    item_id: int,
    payload: schemas.InventoryBalanceUpsert,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_manager_or_admin),
):
    get_item(db, item_id)
    get_location(db, payload.location_id)
    balance = get_or_create_balance(db, item_id, payload.location_id)
    balance.minimum_quantity = payload.minimum_quantity
    balance.par_quantity = payload.par_quantity
    balance.maximum_quantity = payload.maximum_quantity
    db.commit()
    db.refresh(balance)
    return balance


@router.get("/settings/targets", response_model=dict)
def planning_settings(
    location_id: int = Query(gt=0),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    try:
        rows = inventory_settings_rows(db, location_id)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    return {"location_id": location_id, "default_tolerance_percent": 10, "rows": rows}


@router.put("/settings/targets", response_model=dict)
def update_planning_settings(
    payload: schemas.InventoryPlanningSettingsBulkUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_manager_or_admin),
):
    get_location(db, payload.location_id)
    seen_items = set()
    for row in payload.rows:
        if row.inventory_item_id in seen_items:
            raise HTTPException(status_code=400, detail="Each inventory item may appear only once")
        seen_items.add(row.inventory_item_id)
        item = get_item(db, row.inventory_item_id)
        balance = get_or_create_balance(db, item.id, payload.location_id)
        balance.planning_active = row.planning_active
        balance.lower_tolerance_percent = row.lower_tolerance_percent
        balance.upper_tolerance_percent = row.upper_tolerance_percent
        for weekday, target_quantity in row.weekday_targets.items():
            if weekday not in range(7):
                raise HTTPException(status_code=400, detail="Weekday keys must be between 0 and 6")
            target = (
                db.query(InventoryWeekdayTarget)
                .filter(
                    InventoryWeekdayTarget.inventory_item_id == item.id,
                    InventoryWeekdayTarget.location_id == payload.location_id,
                    InventoryWeekdayTarget.weekday == weekday,
                )
                .first()
            )
            if target_quantity is None:
                if target:
                    db.delete(target)
            elif target:
                target.target_quantity = target_quantity
            else:
                db.add(
                    InventoryWeekdayTarget(
                        inventory_item_id=item.id,
                        location_id=payload.location_id,
                        weekday=weekday,
                        target_quantity=target_quantity,
                    )
                )

        vendor_items = (
            db.query(VendorItem).filter(VendorItem.inventory_item_id == item.id).all()
        )
        for vendor_item in vendor_items:
            vendor_item.preferred = False
        if row.preferred_vendor_id is not None:
            vendor = (
                db.query(Vendor)
                .filter(Vendor.id == row.preferred_vendor_id, Vendor.active.is_(True))
                .first()
            )
            if vendor is None:
                raise HTTPException(status_code=404, detail=f"Vendor {row.preferred_vendor_id} not found")
            vendor_item = next(
                (entry for entry in vendor_items if entry.vendor_id == vendor.id), None
            )
            if vendor_item is None:
                vendor_item = VendorItem(
                    vendor_id=vendor.id,
                    inventory_item_id=item.id,
                )
                db.add(vendor_item)
            vendor_item.preferred = True
            vendor_item.vendor_sku = row.vendor_sku
            vendor_item.pack_quantity = row.pack_quantity
            vendor_item.unit_price_cents = row.unit_price_cents
        if row.purchase_unit:
            item.purchase_unit = row.purchase_unit
    db.commit()
    return {
        "location_id": payload.location_id,
        "default_tolerance_percent": 10,
        "rows": inventory_settings_rows(db, payload.location_id),
    }


@router.get("/purchase-order-planner", response_model=dict)
def purchase_order_planner(
    location_id: int = Query(gt=0),
    delivery_date: date = Query(),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    try:
        rows = purchase_order_planner_rows(db, location_id, delivery_date)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    return {
        "location_id": location_id,
        "delivery_date": delivery_date,
        "weekday": delivery_date.weekday(),
        "rows": rows,
    }


def query_stock(
    db: Session,
    *,
    location_id: int | None = None,
    status_filter: str | None = None,
) -> list[schemas.InventoryStockRead]:
    query = (
        db.query(InventoryBalance, InventoryItem, InventoryLocation)
        .join(InventoryItem, InventoryItem.id == InventoryBalance.inventory_item_id)
        .join(InventoryLocation, InventoryLocation.id == InventoryBalance.location_id)
        .filter(InventoryItem.active.is_(True), InventoryLocation.active.is_(True))
    )
    if location_id:
        query = query.filter(InventoryBalance.location_id == location_id)
    rows = query.order_by(InventoryItem.category, InventoryItem.name).all()
    results = []
    for balance, item, location in rows:
        status_value = "OUT" if balance.quantity_on_hand <= 0 else "LOW" if balance.quantity_on_hand <= balance.minimum_quantity else "HEALTHY"
        if status_filter and status_value != status_filter.upper():
            continue
        results.append(schemas.InventoryStockRead(
            inventory_item_id=item.id,
            item_name=item.name,
            category=item.category,
            location_id=location.id,
            location_name=location.name,
            base_unit=item.base_unit,
            quantity_on_hand=balance.quantity_on_hand,
            minimum_quantity=balance.minimum_quantity,
            par_quantity=balance.par_quantity,
            status=status_value,
        ))
    return results


@router.get("/stock", response_model=list[schemas.InventoryStockRead])
def stock(
    location_id: int | None = Query(default=None),
    status_filter: str | None = Query(default=None, alias="status"),
    expiring_days: int | None = Query(default=None, ge=0),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return query_stock(
        db,
        location_id=location_id,
        status_filter=status_filter,
    )


@router.get("/dashboard")
def dashboard(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    balance_rows = (
        db.query(InventoryBalance, InventoryItem, InventoryLocation)
        .join(InventoryItem, InventoryItem.id == InventoryBalance.inventory_item_id)
        .join(InventoryLocation, InventoryLocation.id == InventoryBalance.location_id)
        .filter(InventoryItem.active.is_(True), InventoryLocation.active.is_(True))
        .all()
    )

    def money_cents(quantity: Decimal, unit_cost_cents: int) -> int:
        return int(
            (quantity * Decimal(unit_cost_cents)).quantize(
                Decimal("1"), rounding=ROUND_HALF_UP
            )
        )

    total_inventory_value_cents = 0
    estimated_order_cost_cents = 0
    out_of_stock_count = 0
    category_values: dict[str, dict] = {}
    low_stock_items = []
    balance_count_by_location: dict[int, int] = {}

    for balance, item, location in balance_rows:
        on_hand = Decimal(balance.quantity_on_hand or 0)
        minimum = Decimal(balance.minimum_quantity or 0)
        par = Decimal(balance.par_quantity or 0)
        order_quantity = max(par - on_hand, Decimal("0"))
        inventory_value = money_cents(max(on_hand, Decimal("0")), item.cost_cents)
        order_cost = money_cents(order_quantity, item.cost_cents)
        total_inventory_value_cents += inventory_value
        estimated_order_cost_cents += order_cost
        balance_count_by_location[location.id] = (
            balance_count_by_location.get(location.id, 0) + 1
        )

        category = item.category or "Uncategorized"
        category_row = category_values.setdefault(
            category,
            {"category": category, "value_cents": 0, "item_ids": set()},
        )
        category_row["value_cents"] += inventory_value
        category_row["item_ids"].add(item.id)

        has_stock_target = minimum > 0 or par > 0
        stock_status = (
            "OUT"
            if has_stock_target and on_hand <= 0
            else "LOW"
            if has_stock_target and on_hand <= minimum
            else "HEALTHY"
        )
        if stock_status == "OUT":
            out_of_stock_count += 1
        if stock_status in {"LOW", "OUT"}:
            low_stock_items.append(
                {
                    "inventory_item_id": item.id,
                    "item_name": item.name,
                    "category": item.category,
                    "location_id": location.id,
                    "location_name": location.name,
                    "base_unit": item.base_unit,
                    "quantity_on_hand": on_hand,
                    "minimum_quantity": minimum,
                    "par_quantity": par,
                    "order_quantity": order_quantity,
                    "unit_cost_cents": item.cost_cents,
                    "estimated_cost_cents": order_cost,
                    "status": stock_status,
                }
            )

    category_value_rows = []
    for category_row in category_values.values():
        value_cents = category_row["value_cents"]
        if value_cents <= 0:
            continue
        category_value_rows.append(
            {
                "category": category_row["category"],
                "value_cents": value_cents,
                "item_count": len(category_row["item_ids"]),
                "percent": round(
                    value_cents / total_inventory_value_cents * 100, 1
                )
                if total_inventory_value_cents
                else 0,
            }
        )
    category_value_rows.sort(
        key=lambda row: (-row["value_cents"], row["category"].casefold())
    )
    low_stock_items.sort(
        key=lambda row: (
            row["status"] != "OUT",
            -row["estimated_cost_cents"],
            row["item_name"].casefold(),
        )
    )

    count_rows = (
        db.query(
            InventoryCount.id,
            InventoryCount.location_id,
            InventoryCount.status,
            InventoryCount.updated_at,
        )
        .filter(InventoryCount.status != InventoryCountStatus.REJECTED)
        .order_by(InventoryCount.created_at.desc(), InventoryCount.id.desc())
        .all()
    )
    latest_id_by_location = {}
    open_id_by_location = {}
    open_statuses = {InventoryCountStatus.DRAFT, InventoryCountStatus.SUBMITTED}
    last_count_at = None
    open_count_sessions = 0
    for count_row in count_rows:
        latest_id_by_location.setdefault(count_row.location_id, count_row.id)
        if count_row.status in open_statuses:
            open_count_sessions += 1
            open_id_by_location.setdefault(count_row.location_id, count_row.id)
        if count_row.status == InventoryCountStatus.POSTED and (
            last_count_at is None or count_row.updated_at > last_count_at
        ):
            last_count_at = count_row.updated_at

    selected_count_ids = set(latest_id_by_location.values()) | set(
        open_id_by_location.values()
    )
    selected_counts = (
        db.query(InventoryCount)
        .options(
            joinedload(InventoryCount.location),
            selectinload(InventoryCount.lines).joinedload(InventoryCountLine.item),
        )
        .filter(InventoryCount.id.in_(selected_count_ids))
        .all()
        if selected_count_ids
        else []
    )
    count_by_id = {count.id: count for count in selected_counts}
    latest_by_location = {
        location_id: count_by_id[count_id]
        for location_id, count_id in latest_id_by_location.items()
    }
    open_by_location = {
        location_id: count_by_id[count_id]
        for location_id, count_id in open_id_by_location.items()
    }

    active_locations = (
        db.query(InventoryLocation)
        .filter(InventoryLocation.active.is_(True))
        .order_by(InventoryLocation.name)
        .all()
    )
    count_completion = []
    uncounted_item_count = 0
    counts_requiring_review = 0
    for location in active_locations:
        count = open_by_location.get(location.id)
        if count is None:
            total_lines = balance_count_by_location.get(location.id, 0)
            counted_lines = 0
            exception_count = 0
            status_value = "NOT_STARTED"
            count_id = None
            updated_at = None
        else:
            total_lines = len(count.lines)
            counted_lines = sum(line.is_counted for line in count.lines)
            exception_count = sum(
                line.review_status == "NEEDS_REVIEW" for line in count.lines
            )
            uncounted_item_count += total_lines - counted_lines
            counts_requiring_review += exception_count
            status_value = count.status.value
            count_id = count.id
            updated_at = count.updated_at
        count_completion.append(
            {
                "location_id": location.id,
                "location_name": location.name,
                "count_id": count_id,
                "status": status_value,
                "total_lines": total_lines,
                "counted_lines": counted_lines,
                "uncounted_lines": total_lines - counted_lines,
                "exception_count": exception_count,
                "completion_percent": round(counted_lines / total_lines * 100, 1)
                if total_lines
                else 0,
                "updated_at": updated_at,
            }
        )
    count_completion.sort(
        key=lambda row: (
            row["status"] == "NOT_STARTED",
            row["completion_percent"],
            row["location_name"].casefold(),
        )
    )

    largest_variances = []
    signed_variance_value_cents = 0
    absolute_variance_value_cents = 0
    for count in latest_by_location.values():
        for line in count.lines:
            if not line.is_counted:
                continue
            expected = Decimal(line.expected_quantity or 0)
            counted = Decimal(line.counted_quantity or 0)
            variance = counted - expected
            variance_value_cents = money_cents(variance, line.item.cost_cents)
            signed_variance_value_cents += variance_value_cents
            absolute_variance_value_cents += abs(variance_value_cents)
            if variance == 0:
                continue
            largest_variances.append(
                {
                    "count_id": count.id,
                    "inventory_item_id": line.item.id,
                    "item_name": line.item.name,
                    "category": line.item.category,
                    "location_id": count.location_id,
                    "location_name": count.location.name,
                    "base_unit": line.item.base_unit,
                    "expected_quantity": expected,
                    "counted_quantity": counted,
                    "variance_quantity": variance,
                    "variance_value_cents": variance_value_cents,
                    "absolute_variance_value_cents": abs(variance_value_cents),
                    "source": line.source,
                    "review_status": line.review_status,
                    "updated_at": count.updated_at,
                }
            )
    largest_variances.sort(
        key=lambda row: (
            -row["absolute_variance_value_cents"],
            row["item_name"].casefold(),
        )
    )
    active_item_count = (
        db.query(InventoryItem).filter(InventoryItem.active.is_(True)).count()
    )
    missing_cost_item_count = (
        db.query(InventoryItem)
        .filter(
            InventoryItem.active.is_(True),
            InventoryItem.cost_cents <= 0,
        )
        .count()
    )
    return {
        "item_count": active_item_count,
        "missing_cost_item_count": missing_cost_item_count,
        "location_count": len(active_locations),
        "low_stock_count": len(low_stock_items),
        "out_of_stock_count": out_of_stock_count,
        "open_count_sessions": open_count_sessions,
        "pending_purchase_orders": db.query(PurchaseOrder).filter(PurchaseOrder.status.in_([PurchaseOrderStatus.DRAFT, PurchaseOrderStatus.SUBMITTED, PurchaseOrderStatus.PARTIALLY_RECEIVED])).count(),
        "recent_movements": db.query(StockMovement).filter(StockMovement.inventory_item_id.is_not(None)).order_by(StockMovement.created_at.desc()).limit(10).count(),
        "total_inventory_value_cents": total_inventory_value_cents,
        "estimated_order_cost_cents": estimated_order_cost_cents,
        "signed_variance_value_cents": signed_variance_value_cents,
        "absolute_variance_value_cents": absolute_variance_value_cents,
        "uncounted_item_count": uncounted_item_count,
        "counts_requiring_review": counts_requiring_review,
        "last_count_at": last_count_at,
        "category_values": category_value_rows,
        "count_completion": count_completion,
        "low_stock_items": low_stock_items[:25],
        "largest_variances": largest_variances[:15],
    }


@router.get("/vendors", response_model=list[schemas.VendorRead])
def list_vendors(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return db.query(Vendor).filter(Vendor.active.is_(True)).order_by(Vendor.name).all()


@router.post("/vendors", response_model=schemas.VendorRead, status_code=status.HTTP_201_CREATED)
def create_vendor(payload: schemas.VendorCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_manager_or_admin)):
    vendor = Vendor(**payload.model_dump())
    db.add(vendor)
    db.commit()
    db.refresh(vendor)
    return vendor


@router.post("/vendor-items", response_model=schemas.VendorItemRead, status_code=status.HTTP_201_CREATED)
def create_vendor_item(payload: schemas.VendorItemCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_manager_or_admin)):
    get_item(db, payload.inventory_item_id)
    if not db.query(Vendor).filter(Vendor.id == payload.vendor_id).first():
        raise HTTPException(status_code=404, detail="Vendor not found")
    vendor_item = VendorItem(**payload.model_dump())
    db.add(vendor_item)
    db.commit()
    db.refresh(vendor_item)
    return vendor_item


def post_movement(db: Session, payload: schemas.InventoryMovementCreate, actor_user_id: int | None = None) -> StockMovement:
    item = get_item(db, payload.inventory_item_id)
    get_location(db, payload.location_id)
    if payload.source_event_key:
        existing = db.query(StockMovement).filter(StockMovement.source_event_key == payload.source_event_key).first()
        if existing:
            return existing
    balance = get_or_create_balance(db, payload.inventory_item_id, payload.location_id)
    balance.quantity_on_hand += payload.quantity_change
    legacy_ingredient = db.query(Ingredient).filter(Ingredient.name == item.name).first()
    if not legacy_ingredient:
        legacy_ingredient = Ingredient(name=item.name, unit=item.base_unit, active=True)
        db.add(legacy_ingredient)
        db.flush()
    movement = StockMovement(
        ingredient_id=legacy_ingredient.id,
        inventory_item_id=payload.inventory_item_id,
        location_id=payload.location_id,
        quantity_change=float(payload.quantity_change),
        reason=payload.reason,
        notes=payload.notes,
        source_event_key=payload.source_event_key,
        created_by_user_id=actor_user_id,
        lot_number=payload.lot_number,
        expiration_date=payload.expiration_date,
    )
    db.add(movement)
    db.flush()
    return movement


@router.post("/movements", response_model=schemas.InventoryMovementRead, status_code=status.HTTP_201_CREATED)
def create_movement(payload: schemas.InventoryMovementCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_manager_or_admin)):
    movement = post_movement(db, payload, current_user.id)
    db.commit()
    db.refresh(movement)
    return movement_response(movement)


@router.post("/waste", response_model=schemas.InventoryMovementRead, status_code=status.HTTP_201_CREATED)
def record_waste(payload: schemas.InventoryMovementCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_manager_or_admin)):
    payload.reason = payload.reason or "WASTE"
    movement = post_movement(db, payload, current_user.id)
    db.commit()
    db.refresh(movement)
    return movement_response(movement)


@router.post("/adjustments", response_model=schemas.InventoryMovementRead, status_code=status.HTTP_201_CREATED)
def record_adjustment(payload: schemas.InventoryMovementCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_manager_or_admin)):
    movement = post_movement(db, payload, current_user.id)
    db.commit()
    db.refresh(movement)
    return movement_response(movement)


@router.post("/transfers", response_model=list[schemas.InventoryMovementRead], status_code=status.HTTP_201_CREATED)
def transfer_stock(payload: schemas.InventoryTransferCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_manager_or_admin)):
    if payload.from_location_id == payload.to_location_id:
        raise HTTPException(status_code=400, detail="Transfer locations must be different")
    transfer_key = f"transfer:{current_user.id}:{payload.inventory_item_id}:{payload.from_location_id}:{payload.to_location_id}:{payload.quantity}:{date.today().isoformat()}"
    outbound = post_movement(db, schemas.InventoryMovementCreate(inventory_item_id=payload.inventory_item_id, location_id=payload.from_location_id, quantity_change=-payload.quantity, reason="TRANSFER_OUT", source_event_key=f"{transfer_key}:out", notes=payload.notes), current_user.id)
    inbound = post_movement(db, schemas.InventoryMovementCreate(inventory_item_id=payload.inventory_item_id, location_id=payload.to_location_id, quantity_change=payload.quantity, reason="TRANSFER_IN", source_event_key=f"{transfer_key}:in", notes=payload.notes), current_user.id)
    db.commit()
    db.refresh(outbound)
    db.refresh(inbound)
    return [movement_response(outbound), movement_response(inbound)]


@router.get("/movements", response_model=list[schemas.InventoryMovementRead])
def list_movements(limit: int = Query(default=100, ge=1, le=500), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    movements = db.query(StockMovement).filter(StockMovement.inventory_item_id.is_not(None)).order_by(StockMovement.created_at.desc()).limit(limit).all()
    return [movement_response(movement) for movement in movements]


@router.get("/counts", response_model=list[schemas.InventoryCountRead])
def list_counts(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    counts = db.query(InventoryCount).order_by(InventoryCount.created_at.desc()).all()
    return [{"id": count.id, "location_id": count.location_id, "status": count.status, "counted_by_user_id": count.counted_by_user_id, "reviewed_by_user_id": count.reviewed_by_user_id, "notes": count.notes, "lines": [{"inventory_item_id": line.inventory_item_id, "counted_quantity": line.counted_quantity, "expected_quantity": line.expected_quantity, "notes": line.notes} for line in count.lines]} for count in counts]


@router.get("/count-templates", response_model=list[schemas.CountTemplateRead])
def list_count_templates(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    templates = (
        db.query(InventoryCountTemplate)
        .options(
            selectinload(InventoryCountTemplate.lines).joinedload(
                InventoryCountTemplateLine.item
            ),
            selectinload(InventoryCountTemplate.lines).joinedload(
                InventoryCountTemplateLine.location
            ),
        )
        .filter(InventoryCountTemplate.active.is_(True))
        .order_by(InventoryCountTemplate.name)
        .all()
    )
    return [serialize_template(template) for template in templates]


@router.post(
    "/count-templates",
    response_model=schemas.CountTemplateRead,
    status_code=status.HTTP_201_CREATED,
)
def create_count_template(
    payload: schemas.CountTemplateCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_manager_or_admin),
):
    if (
        db.query(InventoryCountTemplate)
        .filter(func.lower(InventoryCountTemplate.name) == payload.name.strip().lower())
        .first()
    ):
        raise HTTPException(status_code=409, detail="A count template with this name already exists")
    template = InventoryCountTemplate(
        name=payload.name.strip(),
        description=payload.description,
        created_by_user_id=current_user.id,
    )
    db.add(template)
    db.flush()
    for row in payload.lines:
        get_item(db, row.inventory_item_id)
        get_location(db, row.location_id)
        template.lines.append(
            InventoryCountTemplateLine(
                inventory_item_id=row.inventory_item_id,
                location_id=row.location_id,
                display_order=row.display_order,
                preferred_unit=row.preferred_unit,
            )
        )
    db.commit()
    db.refresh(template)
    return serialize_template(template)


def get_count_sheet(db: Session, count_id: int) -> InventoryCount:
    count = (
        db.query(InventoryCount)
        .options(
            joinedload(InventoryCount.location),
            joinedload(InventoryCount.template),
            selectinload(InventoryCount.lines).joinedload(
                InventoryCountLine.item
            ),
        )
        .filter(InventoryCount.id == count_id)
        .first()
    )
    if not count:
        raise HTTPException(status_code=404, detail="Count sheet not found")
    return count


@router.get("/count-sheets", response_model=list[schemas.CountSheetRead])
def list_count_sheets(
    location_id: int | None = Query(default=None),
    count_status: InventoryCountStatus | None = Query(default=None, alias="status"),
    limit: int = Query(default=25, ge=1, le=100),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(InventoryCount).options(
        joinedload(InventoryCount.location),
        joinedload(InventoryCount.template),
        selectinload(InventoryCount.lines).joinedload(InventoryCountLine.item),
    )
    if location_id is not None:
        query = query.filter(InventoryCount.location_id == location_id)
    if count_status is not None:
        query = query.filter(InventoryCount.status == count_status)
    counts = query.order_by(InventoryCount.updated_at.desc()).limit(limit).all()
    return [serialize_count_sheet(db, count) for count in counts]


@router.post(
    "/count-sheets",
    response_model=schemas.CountSheetRead,
    status_code=status.HTTP_201_CREATED,
)
def create_count_sheet(
    payload: schemas.CountSheetCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    get_location(db, payload.location_id)
    if payload.resume_existing:
        existing_query = db.query(InventoryCount).filter(
            InventoryCount.location_id == payload.location_id,
            InventoryCount.status == InventoryCountStatus.DRAFT,
            InventoryCount.counted_by_user_id == current_user.id,
        )
        if payload.template_id is not None:
            existing_query = existing_query.filter(
                InventoryCount.template_id == payload.template_id
            )
        else:
            existing_query = existing_query.filter(
                InventoryCount.template_id.is_(None)
            )
        existing = existing_query.order_by(InventoryCount.updated_at.desc()).first()
        if existing:
            populate_count_sheet(
                db, existing, template_id=existing.template_id
            )
            db.commit()
            db.refresh(existing)
            return serialize_count_sheet(db, existing)
    count = InventoryCount(
        location_id=payload.location_id,
        template_id=payload.template_id,
        counted_by_user_id=current_user.id,
        notes=payload.notes,
        revision=1,
    )
    db.add(count)
    db.flush()
    populate_count_sheet(db, count, template_id=payload.template_id)
    db.commit()
    db.refresh(count)
    return serialize_count_sheet(db, count)


@router.get(
    "/count-sheets/{count_id}",
    response_model=schemas.CountSheetRead,
)
def read_count_sheet(
    count_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return serialize_count_sheet(db, get_count_sheet(db, count_id))


def count_sheet_export_rows(db: Session, count: InventoryCount) -> tuple[dict, list[list]]:
    sheet = serialize_count_sheet(db, count)
    rows = [
        [
            line["display_order"] + 1,
            line["item_name"],
            line["category"] or "",
            line["base_unit"],
            line["expected_quantity"],
            line["par_quantity"],
            line["counted_quantity"] if line["is_counted"] else "",
            line["delta"] if line["is_counted"] else "",
            line["source"] or "",
            line["review_status"],
            line["notes"] or "",
        ]
        for line in sheet["lines"]
    ]
    return sheet, rows


@router.get("/count-sheets/{count_id}/export.csv")
def export_count_sheet_csv(
    count_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    sheet, rows = count_sheet_export_rows(db, get_count_sheet(db, count_id))
    output = StringIO()
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(
        ["Order", "Item", "Category", "Unit", "On Hand", "Par", "Count", "Delta", "Source", "Approval", "Notes"]
    )
    writer.writerows(rows)
    filename = f"inventory-count-{sheet['id']}.csv"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/count-sheets/{count_id}/export.xlsx")
def export_count_sheet_xlsx(
    count_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    sheet, rows = count_sheet_export_rows(db, get_count_sheet(db, count_id))
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Inventory Count"
    worksheet.append([f"{sheet['location_name']} Inventory Count"])
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
    worksheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor="173853")
    worksheet["A1"].alignment = Alignment(horizontal="center")
    headers = ["Order", "Item", "Category", "Unit", "On Hand", "Par", "Count", "Delta", "Source", "Approval", "Notes"]
    worksheet.append(headers)
    for cell in worksheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="295274")
    for row in rows:
        worksheet.append(row)
    worksheet.freeze_panes = "E3"
    worksheet.auto_filter.ref = f"A2:K{worksheet.max_row}"
    widths = [8, 30, 20, 12, 12, 12, 12, 12, 12, 16, 30]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"inventory-count-{sheet['id']}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get(
    "/count-sheets/{count_id}/print",
    response_class=HTMLResponse,
)
def print_count_sheet(
    count_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    sheet, rows = count_sheet_export_rows(db, get_count_sheet(db, count_id))
    body = "".join(
        "<tr>"
        + "".join(f"<td>{escape(str(value))}</td>" for value in row)
        + "</tr>"
        for row in rows
    )
    return HTMLResponse(
        f"""<!doctype html><html><head><meta charset="utf-8"><title>{escape(sheet['location_name'])} Inventory Count</title>
        <style>body{{font:12px Arial,sans-serif;margin:24px;color:#111}}h1{{margin:0 0 5px}}p{{margin:0 0 18px;color:#555}}table{{width:100%;border-collapse:collapse}}th,td{{border:1px solid #bbb;padding:6px;text-align:left}}th{{background:#e8edf2}}@media print{{button{{display:none}}}}</style>
        </head><body><button onclick="window.print()">Print</button><h1>{escape(sheet['location_name'])} Inventory Count</h1>
        <p>Sheet #{sheet['id']} · {escape(sheet['status'])} · {sheet['counted_line_count']} of {sheet['line_count']} counted</p>
        <table><thead><tr>{''.join(f'<th>{escape(header)}</th>' for header in ['Order','Item','Category','Unit','On Hand','Par','Count','Delta','Source','Approval','Notes'])}</tr></thead><tbody>{body}</tbody></table></body></html>"""
    )


@router.patch(
    "/count-sheets/{count_id}/lines/{line_id}",
    response_model=schemas.CountSheetRead,
)
def patch_count_sheet_line(
    count_id: int,
    line_id: int,
    payload: schemas.CountSheetLinePatch,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    count = get_count_sheet(db, count_id)
    line = (
        db.query(InventoryCountLine)
        .filter(
            InventoryCountLine.id == line_id,
            InventoryCountLine.count_id == count.id,
        )
        .first()
    )
    if not line:
        raise HTTPException(status_code=404, detail="Count-sheet row not found")
    apply_line_changes(
        db,
        count,
        line,
        payload.model_dump(exclude_unset=True),
        user_id=current_user.id,
    )
    db.commit()
    db.refresh(count)
    return serialize_count_sheet(db, count)


@router.post(
    "/count-sheets/{count_id}/lines/batch",
    response_model=schemas.CountSheetRead,
)
def patch_count_sheet_lines(
    count_id: int,
    payload: schemas.CountSheetBatchPatch,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    count = get_count_sheet(db, count_id)
    line_ids = [row.line_id for row in payload.edits]
    lines = {
        line.id: line
        for line in db.query(InventoryCountLine)
        .filter(
            InventoryCountLine.count_id == count.id,
            InventoryCountLine.id.in_(line_ids),
        )
        .all()
    }
    if len(lines) != len(set(line_ids)):
        raise HTTPException(status_code=404, detail="One or more count-sheet rows were not found")
    balances = {
        balance.inventory_item_id: balance
        for balance in db.query(InventoryBalance)
        .filter(
            InventoryBalance.location_id == count.location_id,
            InventoryBalance.inventory_item_id.in_(
                [line.inventory_item_id for line in lines.values()]
            ),
        )
        .all()
    }
    for edit in payload.edits:
        changes = edit.model_dump(exclude_unset=True)
        changes.pop("line_id", None)
        apply_line_changes(
            db,
            count,
            lines[edit.line_id],
            changes,
            user_id=current_user.id,
            balance=balances.get(lines[edit.line_id].inventory_item_id),
        )
    db.commit()
    db.refresh(count)
    return serialize_count_sheet(db, count)


@router.post(
    "/count-sheets/{count_id}/reorder",
    response_model=schemas.CountSheetRead,
)
def reorder_count_sheet(
    count_id: int,
    payload: schemas.CountSheetReorder,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    count = get_count_sheet(db, count_id)
    if count.status != InventoryCountStatus.DRAFT:
        raise HTTPException(status_code=409, detail="Only draft count sheets can be reordered")
    if len(payload.line_ids) != len(set(payload.line_ids)):
        raise HTTPException(status_code=422, detail="Count-sheet row order contains duplicates")
    lines = {line.id: line for line in count.lines}
    if set(payload.line_ids) != set(lines):
        raise HTTPException(status_code=422, detail="Count-sheet row order must include every row")
    for display_order, line_id in enumerate(payload.line_ids):
        lines[line_id].display_order = display_order
    count.revision = (count.revision or 0) + 1
    db.commit()
    db.refresh(count)
    return serialize_count_sheet(db, count)


@router.post(
    "/count-sheets/{count_id}/approve",
    response_model=schemas.CountSheetRead,
)
def approve_count_sheet(
    count_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_manager_or_admin),
):
    count = get_count_sheet(db, count_id)
    if count.status != InventoryCountStatus.DRAFT:
        raise HTTPException(status_code=409, detail="Only draft count sheets can be approved")
    ensure_voice_count_ready(db, count.id)
    ensure_count_sheet_complete(count)
    for line in count.lines:
        line.review_status = "APPROVED"
    count.status = InventoryCountStatus.APPROVED
    count.reviewed_by_user_id = current_user.id
    count.approved_at = datetime.utcnow()
    count.revision = (count.revision or 0) + 1
    db.commit()
    db.refresh(count)
    return serialize_count_sheet(db, count)


@router.post("/counts", response_model=schemas.InventoryCountRead, status_code=status.HTTP_201_CREATED)
def create_count(payload: schemas.InventoryCountCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    get_location(db, payload.location_id)
    count = InventoryCount(location_id=payload.location_id, counted_by_user_id=current_user.id, notes=payload.notes)
    db.add(count)
    db.flush()
    for line in payload.lines:
        get_item(db, line.inventory_item_id)
        balance = get_or_create_balance(db, line.inventory_item_id, payload.location_id)
        db.add(InventoryCountLine(count_id=count.id, inventory_item_id=line.inventory_item_id, counted_quantity=line.counted_quantity, expected_quantity=balance.quantity_on_hand, notes=line.notes, is_counted=True, source="MANUAL", review_status="READY"))
    db.commit()
    db.refresh(count)
    return {"id": count.id, "location_id": count.location_id, "status": count.status, "counted_by_user_id": count.counted_by_user_id, "reviewed_by_user_id": count.reviewed_by_user_id, "notes": count.notes, "lines": [{"inventory_item_id": line.inventory_item_id, "counted_quantity": line.counted_quantity, "expected_quantity": line.expected_quantity, "notes": line.notes} for line in count.lines]}


@router.post("/counts/{count_id}/submit", response_model=schemas.InventoryCountRead)
def submit_count(count_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    count = db.query(InventoryCount).filter(InventoryCount.id == count_id).first()
    if not count:
        raise HTTPException(status_code=404, detail="Count session not found")
    if count.status != InventoryCountStatus.DRAFT:
        raise HTTPException(status_code=400, detail="Only draft counts can be submitted")
    ensure_voice_count_ready(db, count.id)
    ensure_count_sheet_complete(count)
    count.status = InventoryCountStatus.SUBMITTED
    db.commit()
    return {"id": count.id, "location_id": count.location_id, "status": count.status, "counted_by_user_id": count.counted_by_user_id, "reviewed_by_user_id": count.reviewed_by_user_id, "notes": count.notes, "lines": []}


@router.post("/counts/{count_id}/post", response_model=schemas.InventoryCountRead)
def post_count(count_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_manager_or_admin)):
    count = db.query(InventoryCount).filter(InventoryCount.id == count_id).first()
    if not count:
        raise HTTPException(status_code=404, detail="Count session not found")
    if count.status not in {InventoryCountStatus.SUBMITTED, InventoryCountStatus.APPROVED}:
        raise HTTPException(status_code=400, detail="Count must be submitted before posting")
    ensure_voice_count_ready(db, count.id)
    ensure_count_sheet_complete(count)
    for line in count.lines:
        balance = get_or_create_balance(db, line.inventory_item_id, count.location_id)
        difference = line.counted_quantity - balance.quantity_on_hand
        if difference:
            post_movement(db, schemas.InventoryMovementCreate(inventory_item_id=line.inventory_item_id, location_id=count.location_id, quantity_change=difference, reason="COUNT", source_event_key=f"count:{count.id}:item:{line.inventory_item_id}"), current_user.id)
    count.status = InventoryCountStatus.POSTED
    count.reviewed_by_user_id = current_user.id
    db.commit()
    return {"id": count.id, "location_id": count.location_id, "status": count.status, "counted_by_user_id": count.counted_by_user_id, "reviewed_by_user_id": count.reviewed_by_user_id, "notes": count.notes, "lines": []}


def enrich_csv_preview(
    db: Session,
    preview: dict,
    payload: schemas.PurchaseOrderCSVPreviewRequest,
) -> dict:
    location_ids = set(payload.location_overrides.values())
    if payload.default_location_id:
        location_ids.add(payload.default_location_id)
    locations = {
        location.id: location
        for location in db.query(InventoryLocation)
        .filter(InventoryLocation.id.in_(location_ids or {-1}), InventoryLocation.active.is_(True))
        .all()
    }
    missing_locations = location_ids - set(locations)
    if missing_locations:
        raise HTTPException(
            status_code=404,
            detail=f"Inventory location {sorted(missing_locations)[0]} not found",
        )
    planner_by_location = {}
    if payload.expected_date:
        for location_id in location_ids:
            planner_by_location[location_id] = {
                row["inventory_item_id"]: row
                for row in purchase_order_planner_rows(
                    db, location_id, payload.expected_date
                )
            }
    for row in preview["rows"]:
        location_id = payload.location_overrides.get(
            row["row_number"], payload.default_location_id
        )
        row["location_id"] = location_id
        row["location_name"] = locations[location_id].name if location_id else None
        planning = (
            planner_by_location.get(location_id, {}).get(row["inventory_item_id"])
            if row["inventory_item_id"] and location_id
            else None
        )
        row["planning"] = planning
    return preview


@router.post("/purchase-orders/import-preview", response_model=dict)
def preview_purchase_order_import(
    payload: schemas.PurchaseOrderCSVPreviewRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_manager_or_admin),
):
    if not db.query(Vendor).filter(Vendor.id == payload.vendor_id).first():
        raise HTTPException(status_code=404, detail="Vendor not found")
    try:
        preview = preview_purchase_order_csv(
            db,
            vendor_id=payload.vendor_id,
            csv_text=payload.csv_text,
            item_overrides=payload.item_overrides,
        )
    except PurchaseOrderCSVError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    preview = enrich_csv_preview(db, preview, payload)
    existing = (
        db.query(PurchaseOrder)
        .filter(PurchaseOrder.import_source_hash == preview["source_hash"])
        .first()
    )
    preview["duplicate_order_id"] = existing.id if existing else None
    return preview


@router.post(
    "/purchase-orders/import-csv",
    response_model=schemas.PurchaseOrderRead,
    status_code=status.HTTP_201_CREATED,
)
def import_purchase_order_csv(
    payload: schemas.PurchaseOrderCSVImportRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_manager_or_admin),
):
    if not db.query(Vendor).filter(Vendor.id == payload.vendor_id).first():
        raise HTTPException(status_code=404, detail="Vendor not found")
    try:
        preview = preview_purchase_order_csv(
            db,
            vendor_id=payload.vendor_id,
            csv_text=payload.csv_text,
            item_overrides=payload.item_overrides,
        )
    except PurchaseOrderCSVError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    preview = enrich_csv_preview(db, preview, payload)
    duplicate = (
        db.query(PurchaseOrder)
        .filter(PurchaseOrder.import_source_hash == preview["source_hash"])
        .first()
    )
    if duplicate:
        raise HTTPException(
            status_code=409,
            detail=f"This CSV was already imported as purchase order #{duplicate.id}.",
        )
    if not preview["ready_to_import"]:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Resolve all CSV rows before importing: {preview['unresolved_count']} unmatched and "
                f"{preview['invalid_count']} invalid."
            ),
        )
    external_reference = (payload.external_reference or preview["external_reference"] or "").strip() or None
    if external_reference:
        existing_reference = (
            db.query(PurchaseOrder)
            .filter(
                PurchaseOrder.vendor_id == payload.vendor_id,
                PurchaseOrder.external_reference == external_reference,
            )
            .first()
        )
        if existing_reference:
            raise HTTPException(
                status_code=409,
                detail=(
                    f"Vendor reference {external_reference} is already used by purchase order "
                    f"#{existing_reference.id}."
                ),
            )
    filename = (payload.source_filename or "").replace("\\", "/").rsplit("/", 1)[-1].strip() or None
    purchase_order = PurchaseOrder(
        vendor_id=payload.vendor_id,
        expected_date=payload.expected_date,
        notes=payload.notes,
        external_reference=external_reference,
        import_source_hash=preview["source_hash"],
        imported_filename=filename,
        created_by_user_id=current_user.id,
    )
    db.add(purchase_order)
    db.flush()
    for row in preview["rows"]:
        item = get_item(db, row["inventory_item_id"])
        vendor_item = (
            db.query(VendorItem)
            .filter(
                VendorItem.vendor_id == payload.vendor_id,
                VendorItem.inventory_item_id == item.id,
            )
            .first()
        )
        pack_quantity = Decimal(
            (vendor_item.pack_quantity if vendor_item else item.purchase_to_base) or 1
        )
        purchase_order.lines.append(
            PurchaseOrderLine(
                inventory_item_id=row["inventory_item_id"],
                location_id=row["location_id"],
                ordered_quantity=Decimal(row["quantity"]),
                unit_price_cents=row["unit_price_cents"],
                purchase_unit=item.purchase_unit or item.base_unit,
                quantity_per_purchase_unit=pack_quantity,
            )
        )
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="This purchase order CSV or vendor reference was imported by another request.",
        ) from exc
    db.refresh(purchase_order)
    return purchase_order_response(purchase_order)


@router.post(
    "/purchase-orders/from-plan",
    response_model=list[schemas.PurchaseOrderRead],
    status_code=status.HTTP_201_CREATED,
)
def create_purchase_orders_from_plan(
    payload: schemas.PurchaseOrderPlanCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_manager_or_admin),
):
    grouped: dict[int, list[tuple[schemas.PurchaseOrderPlanLineCreate, InventoryItem, VendorItem]]] = {}
    seen = set()
    missing_vendors = []
    for line in payload.lines:
        key = (line.inventory_item_id, line.location_id)
        if key in seen:
            raise HTTPException(status_code=400, detail="Each item and area may appear only once")
        seen.add(key)
        item = get_item(db, line.inventory_item_id)
        get_location(db, line.location_id)
        balance = (
            db.query(InventoryBalance)
            .filter(
                InventoryBalance.inventory_item_id == item.id,
                InventoryBalance.location_id == line.location_id,
                InventoryBalance.planning_active.is_(True),
            )
            .first()
        )
        if balance is None:
            raise HTTPException(
                status_code=400,
                detail=f"{item.name} is not enabled for ordering in this area",
            )
        vendor_item = (
            db.query(VendorItem)
            .join(Vendor, Vendor.id == VendorItem.vendor_id)
            .filter(
                VendorItem.inventory_item_id == item.id,
                VendorItem.preferred.is_(True),
                Vendor.active.is_(True),
            )
            .first()
        )
        if vendor_item is None:
            missing_vendors.append(item.name)
            continue
        grouped.setdefault(vendor_item.vendor_id, []).append((line, item, vendor_item))
    if missing_vendors:
        raise HTTPException(
            status_code=409,
            detail="Assign a preferred vendor before ordering: " + ", ".join(missing_vendors[:8]),
        )

    orders = []
    for vendor_id, rows in grouped.items():
        order = PurchaseOrder(
            vendor_id=vendor_id,
            expected_date=payload.expected_date,
            notes=payload.notes,
            created_by_user_id=current_user.id,
        )
        db.add(order)
        db.flush()
        for line, item, vendor_item in rows:
            pack_quantity = Decimal(vendor_item.pack_quantity or item.purchase_to_base or 1)
            unit_price_cents = vendor_item.unit_price_cents or int(
                Decimal(item.cost_cents or 0) * pack_quantity
            )
            order.lines.append(
                PurchaseOrderLine(
                    inventory_item_id=item.id,
                    location_id=line.location_id,
                    ordered_quantity=line.purchase_quantity,
                    unit_price_cents=unit_price_cents,
                    purchase_unit=item.purchase_unit or item.base_unit,
                    quantity_per_purchase_unit=pack_quantity,
                )
            )
        orders.append(order)
    db.commit()
    for order in orders:
        db.refresh(order)
    return [purchase_order_response(order) for order in orders]


@router.post("/purchase-orders", response_model=schemas.PurchaseOrderRead, status_code=status.HTTP_201_CREATED)
def create_purchase_order(payload: schemas.PurchaseOrderCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_manager_or_admin)):
    if not db.query(Vendor).filter(Vendor.id == payload.vendor_id).first():
        raise HTTPException(status_code=404, detail="Vendor not found")
    purchase_order = PurchaseOrder(vendor_id=payload.vendor_id, expected_date=payload.expected_date, notes=payload.notes, created_by_user_id=current_user.id)
    db.add(purchase_order)
    db.flush()
    for line in payload.lines:
        item = get_item(db, line.inventory_item_id)
        if line.location_id:
            get_location(db, line.location_id)
        purchase_order.lines.append(PurchaseOrderLine(
            inventory_item_id=line.inventory_item_id,
            location_id=line.location_id,
            ordered_quantity=line.ordered_quantity,
            unit_price_cents=line.unit_price_cents,
            purchase_unit=line.purchase_unit or item.purchase_unit or item.base_unit,
            quantity_per_purchase_unit=line.quantity_per_purchase_unit,
        ))
    db.commit()
    db.refresh(purchase_order)
    return purchase_order_response(purchase_order)


@router.get("/purchase-orders", response_model=list[schemas.PurchaseOrderRead])
def list_purchase_orders(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    orders = db.query(PurchaseOrder).order_by(PurchaseOrder.created_at.desc()).all()
    return [purchase_order_response(order) for order in orders]


@router.put("/purchase-orders/{order_id}", response_model=schemas.PurchaseOrderRead)
def update_draft_purchase_order(
    order_id: int,
    payload: schemas.PurchaseOrderCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_manager_or_admin),
):
    order = db.query(PurchaseOrder).filter(PurchaseOrder.id == order_id).first()
    if order is None:
        raise HTTPException(status_code=404, detail="Purchase order not found")
    if order.status != PurchaseOrderStatus.DRAFT:
        raise HTTPException(status_code=400, detail="Only draft purchase orders can be edited")
    if not db.query(Vendor).filter(Vendor.id == payload.vendor_id).first():
        raise HTTPException(status_code=404, detail="Vendor not found")
    order.vendor_id = payload.vendor_id
    order.expected_date = payload.expected_date
    order.notes = payload.notes
    order.lines.clear()
    db.flush()
    for line in payload.lines:
        item = get_item(db, line.inventory_item_id)
        if line.location_id:
            get_location(db, line.location_id)
        order.lines.append(
            PurchaseOrderLine(
                inventory_item_id=item.id,
                location_id=line.location_id,
                ordered_quantity=line.ordered_quantity,
                unit_price_cents=line.unit_price_cents,
                purchase_unit=line.purchase_unit or item.purchase_unit or item.base_unit,
                quantity_per_purchase_unit=line.quantity_per_purchase_unit,
            )
        )
    if not order.lines:
        raise HTTPException(status_code=400, detail="Purchase order needs at least one line")
    db.commit()
    db.refresh(order)
    return purchase_order_response(order)


@router.post("/purchase-orders/{order_id}/submit", response_model=dict)
def submit_purchase_order(order_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_manager_or_admin)):
    order = db.query(PurchaseOrder).filter(PurchaseOrder.id == order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Purchase order not found")
    if not order.lines:
        raise HTTPException(status_code=400, detail="Purchase order needs at least one line")
    if order.status != PurchaseOrderStatus.DRAFT:
        raise HTTPException(status_code=400, detail="Only draft orders can be submitted")
    order.status = PurchaseOrderStatus.SUBMITTED
    db.commit()
    return {"id": order.id, "status": order.status}


@router.post("/purchase-orders/{order_id}/cancel", response_model=dict)
def cancel_purchase_order(order_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_manager_or_admin)):
    order = db.query(PurchaseOrder).filter(PurchaseOrder.id == order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Purchase order not found")
    if order.status in {PurchaseOrderStatus.RECEIVED, PurchaseOrderStatus.CANCELLED}:
        raise HTTPException(status_code=400, detail="This order cannot be cancelled")
    order.status = PurchaseOrderStatus.CANCELLED
    db.commit()
    return {"id": order.id, "status": order.status}


@router.post("/receiving", response_model=dict, status_code=status.HTTP_201_CREATED)
def receive_purchase_order(payload: schemas.ReceivingCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_manager_or_admin)):
    order = db.query(PurchaseOrder).filter(PurchaseOrder.id == payload.purchase_order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Purchase order not found")
    if order.status not in {PurchaseOrderStatus.SUBMITTED, PurchaseOrderStatus.PARTIALLY_RECEIVED}:
        raise HTTPException(status_code=400, detail="Submit the purchase order before receiving it")
    if not payload.lines:
        raise HTTPException(status_code=400, detail="At least one receiving line is required")
    receiving = InventoryReceiving(purchase_order_id=order.id, received_by_user_id=current_user.id, invoice_number=payload.invoice_number, notes=payload.notes)
    db.add(receiving)
    db.flush()
    for line in payload.lines:
        item = get_item(db, line.inventory_item_id)
        matching_line = None
        if line.purchase_order_line_id is not None:
            matching_line = next(
                (candidate for candidate in order.lines if candidate.id == line.purchase_order_line_id),
                None,
            )
        if matching_line is None:
            candidates = [
                candidate
                for candidate in order.lines
                if candidate.inventory_item_id == line.inventory_item_id
                and (line.location_id is None or candidate.location_id == line.location_id)
            ]
            if len(candidates) == 1:
                matching_line = candidates[0]
        if matching_line is None or matching_line.inventory_item_id != item.id:
            raise HTTPException(status_code=400, detail=f"{item.name} is not a unique line on this order")
        location_id = line.location_id or matching_line.location_id
        if location_id is None:
            raise HTTPException(status_code=400, detail=f"Choose a receiving area for {item.name}")
        get_location(db, location_id)
        remaining = max(
            Decimal(matching_line.ordered_quantity or 0)
            - Decimal(matching_line.received_quantity or 0),
            Decimal("0"),
        )
        if line.received_quantity > remaining and not payload.allow_overage:
            raise HTTPException(
                status_code=400,
                detail=f"Received quantity for {item.name} exceeds the remaining {remaining}",
            )
        conversion = Decimal(matching_line.quantity_per_purchase_unit or 1)
        base_quantity = line.received_quantity * conversion
        receiving.lines.append(
            InventoryReceivingLine(
                purchase_order_line_id=matching_line.id,
                inventory_item_id=item.id,
                location_id=location_id,
                received_quantity=line.received_quantity,
                unit_price_cents=line.unit_price_cents,
                lot_number=line.lot_number,
                expiration_date=line.expiration_date,
                notes=line.notes,
            )
        )
        post_movement(
            db,
            schemas.InventoryMovementCreate(
                inventory_item_id=item.id,
                location_id=location_id,
                quantity_change=base_quantity,
                reason="RECEIVE",
                source_event_key=f"receiving:{receiving.id}:po-line:{matching_line.id}",
                lot_number=line.lot_number,
                expiration_date=line.expiration_date,
                notes=line.notes,
            ),
            current_user.id,
        )
        matching_line.received_quantity += line.received_quantity
    order.status = PurchaseOrderStatus.RECEIVED if all(line.received_quantity >= line.ordered_quantity for line in order.lines) else PurchaseOrderStatus.PARTIALLY_RECEIVED
    db.commit()
    return {
        "id": receiving.id,
        "purchase_order_id": order.id,
        "status": order.status,
        "lines": [
            {
                "purchase_order_line_id": line.purchase_order_line_id,
                "inventory_item_id": line.inventory_item_id,
                "location_id": line.location_id,
                "received_quantity": line.received_quantity,
            }
            for line in receiving.lines
        ],
    }
