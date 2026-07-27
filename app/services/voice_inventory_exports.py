import csv
import html
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models import (
    InventoryVoiceEntry,
    InventoryVoiceReviewStatus,
    InventoryVoiceSession,
)
from app.services.voice_inventory import effective_counts


NAVY = "17324D"
BLUE = "2D6CDF"
PALE_BLUE = "EAF2FF"
PALE_RED = "FDECEC"
PALE_GREEN = "E8F7F0"
WHITE = "FFFFFF"
MUTED = "5B6B7A"
THIN_GRAY = Side(style="thin", color="D9E2EC")


def audit_rows(db: Session, session: InventoryVoiceSession) -> list[dict]:
    entries = (
        db.query(InventoryVoiceEntry)
        .filter(InventoryVoiceEntry.session_id == session.id)
        .order_by(InventoryVoiceEntry.id)
        .all()
    )
    return [
        {
            "timestamp": entry.created_at,
            "location": entry.location.name,
            "transcript": entry.utterance.transcript,
            "evidence": entry.evidence,
            "action": entry.action.value,
            "spoken_item": entry.spoken_item or "",
            "matched_item": entry.item.name if entry.item else "",
            "spoken_quantity": entry.spoken_quantity,
            "spoken_unit": entry.spoken_unit or "",
            "base_quantity": entry.normalized_quantity,
            "base_unit": entry.item.base_unit if entry.item else "",
            "review_status": entry.review_status.value,
            "issue": entry.ambiguity_reason or "",
            "source_entry": entry.supersedes_entry_id or "",
        }
        for entry in entries
    ]


def csv_bytes(db: Session, session: InventoryVoiceSession) -> bytes:
    output = io.StringIO(newline="")
    fieldnames = [
        "timestamp",
        "location",
        "transcript",
        "evidence",
        "action",
        "spoken_item",
        "matched_item",
        "spoken_quantity",
        "spoken_unit",
        "base_quantity",
        "base_unit",
        "review_status",
        "issue",
        "source_entry",
    ]
    writer = csv.DictWriter(output, fieldnames=fieldnames)
    writer.writeheader()
    for row in audit_rows(db, session):
        safe_row = {
            key: (
                value.isoformat()
                if isinstance(value, datetime)
                else str(value)
                if value is not None
                else ""
            )
            for key, value in row.items()
        }
        writer.writerow(safe_row)
    return output.getvalue().encode("utf-8-sig")


def _style_sheet(sheet, title: str, headers: list[str], row_count: int) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = sheet.cell(1, 1, title)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.font = Font(color=WHITE, bold=True, size=16)
    title_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28
    for cell in sheet[3]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)
    if row_count:
        sheet.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{row_count + 3}"
    for row in sheet.iter_rows(min_row=4, max_row=row_count + 3):
        for cell in row:
            cell.border = Border(bottom=THIN_GRAY)
            cell.alignment = Alignment(vertical="top")


def _fit_columns(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def xlsx_bytes(db: Session, session: InventoryVoiceSession) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Count Summary"
    summary_headers = [
        "Location",
        "Inventory Item",
        "Quantity",
        "Base Unit",
        "Source Entry IDs",
    ]
    summary.append([])
    summary.append(
        [
            f"Session #{session.id}",
            f"Manager: {session.manager.full_name}",
            f"Started: {session.started_at:%Y-%m-%d %H:%M}",
            f"Status: {session.status.value}",
            "",
        ]
    )
    summary.append(summary_headers)
    counts = effective_counts(db, session.id)
    for row in counts:
        summary.append(
            [
                row["location_name"],
                row["item_name"],
                float(row["quantity"]),
                row["base_unit"],
                ", ".join(str(value) for value in row["source_entry_ids"]),
            ]
        )
    _style_sheet(summary, "Voice Inventory Count Summary", summary_headers, len(counts))
    _fit_columns(summary, [24, 34, 14, 14, 24])
    for cell in summary["C"][3:]:
        cell.number_format = "#,##0.0000"
        cell.alignment = Alignment(horizontal="right")
    summary.page_setup.orientation = "landscape"
    summary.print_title_rows = "1:3"

    audit = workbook.create_sheet("Voice Audit")
    audit_headers = [
        "Timestamp",
        "Location",
        "Transcript",
        "Evidence",
        "Action",
        "Spoken Item",
        "Matched Item",
        "Spoken Qty",
        "Spoken Unit",
        "Base Qty",
        "Base Unit",
        "Review Status",
        "Issue",
        "Supersedes",
    ]
    audit.append([])
    audit.append(["Every voice-derived action and correction is preserved."] + [""] * 13)
    audit.append(audit_headers)
    rows = audit_rows(db, session)
    for row in rows:
        audit.append(
            [
                row["timestamp"],
                row["location"],
                row["transcript"],
                row["evidence"],
                row["action"],
                row["spoken_item"],
                row["matched_item"],
                float(row["spoken_quantity"]) if row["spoken_quantity"] is not None else None,
                row["spoken_unit"],
                float(row["base_quantity"]) if row["base_quantity"] is not None else None,
                row["base_unit"],
                row["review_status"],
                row["issue"],
                row["source_entry"],
            ]
        )
    _style_sheet(audit, "Voice Inventory Audit", audit_headers, len(rows))
    _fit_columns(audit, [20, 22, 42, 34, 13, 25, 28, 13, 14, 13, 13, 18, 36, 12])
    for row_number in range(4, len(rows) + 4):
        audit.cell(row_number, 1).number_format = "yyyy-mm-dd hh:mm:ss"
        for column in (3, 4, 13):
            audit.cell(row_number, column).alignment = Alignment(
                vertical="top", wrap_text=True
            )
        status = audit.cell(row_number, 12)
        status.fill = PatternFill(
            "solid",
            fgColor=(
                PALE_RED
                if status.value in {"NEEDS_REVIEW", "REJECTED"}
                else PALE_GREEN
            ),
        )
    audit.page_setup.orientation = "landscape"
    audit.print_title_rows = "1:3"

    review = workbook.create_sheet("Needs Review")
    review_headers = [
        "Timestamp",
        "Location",
        "Evidence",
        "Spoken Item",
        "Matched Item",
        "Quantity",
        "Unit",
        "Status",
        "Issue",
    ]
    review.append([])
    review.append(["Entries requiring attention or changed during review."] + [""] * 8)
    review.append(review_headers)
    review_rows = [
        row
        for row in rows
        if row["review_status"]
        in {
            InventoryVoiceReviewStatus.NEEDS_REVIEW.value,
            InventoryVoiceReviewStatus.CORRECTED.value,
            InventoryVoiceReviewStatus.REJECTED.value,
        }
    ]
    for row in review_rows:
        review.append(
            [
                row["timestamp"],
                row["location"],
                row["evidence"],
                row["spoken_item"],
                row["matched_item"],
                float(row["base_quantity"]) if row["base_quantity"] is not None else None,
                row["base_unit"] or row["spoken_unit"],
                row["review_status"],
                row["issue"],
            ]
        )
    _style_sheet(review, "Voice Inventory Review Queue", review_headers, len(review_rows))
    _fit_columns(review, [20, 22, 36, 25, 28, 14, 14, 18, 40])
    for row_number in range(4, len(review_rows) + 4):
        review.cell(row_number, 1).number_format = "yyyy-mm-dd hh:mm:ss"
        review.cell(row_number, 9).alignment = Alignment(
            vertical="top", wrap_text=True
        )

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def print_html(db: Session, session: InventoryVoiceSession) -> str:
    counts = effective_counts(db, session.id)
    by_location: dict[str, list[dict]] = {}
    for row in counts:
        by_location.setdefault(row["location_name"], []).append(row)
    sections = []
    for location, rows in by_location.items():
        table_rows = "".join(
            "<tr>"
            f"<td>{html.escape(row['item_name'])}</td>"
            f"<td class='number'>{row['quantity']}</td>"
            f"<td>{html.escape(row['base_unit'])}</td>"
            "<td class='write'></td>"
            "</tr>"
            for row in rows
        )
        sections.append(
            f"<section><h2>{html.escape(location)}</h2>"
            "<table><thead><tr><th>Item</th><th>Count</th><th>Unit</th>"
            "<th>Manager check</th></tr></thead>"
            f"<tbody>{table_rows}</tbody></table></section>"
        )
    issues = [
        row
        for row in audit_rows(db, session)
        if row["review_status"]
        in {
            InventoryVoiceReviewStatus.NEEDS_REVIEW.value,
            InventoryVoiceReviewStatus.CORRECTED.value,
            InventoryVoiceReviewStatus.REJECTED.value,
        }
    ]
    issue_rows = "".join(
        "<tr>"
        f"<td>{html.escape(str(row['location']))}</td>"
        f"<td>{html.escape(str(row['evidence']))}</td>"
        f"<td>{html.escape(str(row['review_status']))}</td>"
        f"<td>{html.escape(str(row['issue']))}</td>"
        "</tr>"
        for row in issues
    )
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Voice Inventory #{session.id}</title>
<style>
body{{font:12px Arial,sans-serif;color:#172b3a;margin:24px}}h1{{margin:0}}.meta{{color:#5b6b7a;margin:6px 0 22px}}
section{{break-after:page;margin-bottom:24px}}table{{border-collapse:collapse;width:100%}}th{{background:#17324d;color:white}}
th,td{{padding:8px;border-bottom:1px solid #ccd6df;text-align:left}}.number{{text-align:right}}.write{{width:28%;height:28px}}
.sign{{margin-top:35px;border-top:1px solid #333;width:45%;padding-top:6px}}@media print{{body{{margin:.35in}}.no-print{{display:none}}}}
</style></head><body>
<button class="no-print" onclick="window.print()">Print</button>
<h1>Voice Inventory Count</h1>
<div class="meta">Session #{session.id} · {html.escape(session.manager.full_name)} ·
{session.started_at:%Y-%m-%d %H:%M} · {html.escape(session.status.value)}</div>
{''.join(sections) or '<p>No accepted counts.</p>'}
<section><h2>Review and correction audit</h2>
<table><thead><tr><th>Location</th><th>Spoken evidence</th><th>Status</th><th>Issue</th></tr></thead>
<tbody>{issue_rows or '<tr><td colspan="4">No exceptions.</td></tr>'}</tbody></table>
<div class="sign">Manager signature / date</div></section>
</body></html>"""
