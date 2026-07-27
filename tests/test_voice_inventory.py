from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import UserRole
from app.services.voice_inventory import split_inventory_phrases


async def login_manager(client, email: str):
    await client.post(
        "/auth/register",
        json={
            "email": email,
            "password": "secret123",
            "full_name": "Voice Inventory Manager",
            "role": UserRole.MANAGER.value,
        },
    )
    response = await client.post(
        "/auth/login", json={"email": email, "password": "secret123"}
    )
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def test_voice_transition_words_split_entries_without_splitting_item_commas():
    assert split_inventory_phrases(
        "Chicken breast, two cases next tomatoes five pounds "
        "bump onions three bags then carrots four pounds"
    ) == [
        "Chicken breast, two cases",
        "tomatoes five pounds",
        "onions three bags",
        "carrots four pounds",
    ]
    assert split_inventory_phrases("Next Level Sauce two bottles") == [
        "Next Level Sauce two bottles"
    ]


@pytest.mark.asyncio
async def test_voice_transition_words_create_multiple_inventory_entries(client):
    headers = await login_manager(client, "voice-transitions@example.com")
    location = await client.post(
        "/inventory/locations",
        json={"name": "Transition Walk In"},
        headers=headers,
    )
    item_ids = {}
    for name in [
        "Transition Apples",
        "Transition Tomatoes",
        "Transition Onions",
        "Transition Carrots",
    ]:
        response = await client.post(
            "/inventory/items",
            json={"name": name, "base_unit": "each"},
            headers=headers,
        )
        assert response.status_code == 201, response.text
        item_ids[name] = response.json()["id"]
    session = await client.post(
        "/inventory/voice/sessions",
        json={
            "client_session_id": "91111111-1111-4111-8111-111111111111",
            "initial_location_id": location.json()["id"],
        },
        headers=headers,
    )
    result = await client.post(
        f"/inventory/voice/sessions/{session.json()['id']}/utterances",
        json={
            "client_event_id": "92111111-1111-4111-8111-111111111111",
            "sequence": 1,
            "transcript": (
                "Transition Apples, two each next Transition Tomatoes five each "
                "bump Transition Onions three each then Transition Carrots four each"
            ),
        },
        headers=headers,
    )
    assert result.status_code == 201, result.text
    body = result.json()
    assert body["status"] == "ACCEPTED"
    assert len(body["entries"]) == 4
    assert {
        entry["item_name"]: entry["normalized_quantity"]
        for entry in body["entries"]
    } == {
        "Transition Apples": "2.0000",
        "Transition Tomatoes": "5.0000",
        "Transition Onions": "3.0000",
        "Transition Carrots": "4.0000",
    }


@pytest.mark.asyncio
async def test_voice_inventory_multilocation_review_drafts_and_exports(client):
    headers = await login_manager(client, "voice-inventory@example.com")
    walk_in = await client.post(
        "/inventory/locations", json={"name": "Voice Walk In"}, headers=headers
    )
    dry = await client.post(
        "/inventory/locations", json={"name": "Voice Dry Storage"}, headers=headers
    )
    chicken = await client.post(
        "/inventory/items",
        json={
            "name": "Voice Chicken Breast",
            "base_unit": "lb",
            "purchase_unit": "case",
            "purchase_to_base": 40,
        },
        headers=headers,
    )
    tomatoes = await client.post(
        "/inventory/items",
        json={"name": "Voice Tomatoes", "base_unit": "lb"},
        headers=headers,
    )
    assert walk_in.status_code == dry.status_code == chicken.status_code == tomatoes.status_code == 201

    started = await client.post(
        "/inventory/voice/sessions",
        json={
            "client_session_id": "11111111-1111-4111-8111-111111111111",
            "initial_location_id": walk_in.json()["id"],
            "device_metadata": {"browser": "pytest"},
        },
        headers=headers,
    )
    assert started.status_code == 201, started.text
    session_id = started.json()["id"]

    first_payload = {
        "client_event_id": "21111111-1111-4111-8111-111111111111",
        "sequence": 1,
        "transcript": "Voice Chicken Breast 2 cases",
    }
    first = await client.post(
        f"/inventory/voice/sessions/{session_id}/utterances",
        json=first_payload,
        headers=headers,
    )
    duplicate = await client.post(
        f"/inventory/voice/sessions/{session_id}/utterances",
        json=first_payload,
        headers=headers,
    )
    assert first.status_code == duplicate.status_code == 201
    assert duplicate.json()["id"] == first.json()["id"]
    assert first.json()["entries"][0]["normalized_quantity"] == "80.0000"

    repeated = await client.post(
        f"/inventory/voice/sessions/{session_id}/utterances",
        json={
            "client_event_id": "31111111-1111-4111-8111-111111111111",
            "sequence": 2,
            "transcript": "Voice Chicken Breast 3 cases",
        },
        headers=headers,
    )
    assert repeated.status_code == 201, repeated.text
    assert repeated.json()["status"] == "NEEDS_CLARIFICATION"
    unresolved_entry_id = repeated.json()["entries"][0]["id"]

    correction = await client.patch(
        f"/inventory/voice/sessions/{session_id}/entries/{unresolved_entry_id}",
        json={
            "inventory_item_id": chicken.json()["id"],
            "quantity": 3,
            "unit": "case",
        },
        headers=headers,
    )
    assert correction.status_code == 200, correction.text
    assert correction.json()["review_status"] == "CORRECTED"
    assert correction.json()["normalized_quantity"] == "120.0000"
    assert correction.json()["supersedes_entry_id"] == unresolved_entry_id
    before_finish = await client.get(
        f"/inventory/voice/sessions/{session_id}", headers=headers
    )
    assert before_finish.json()["draft_counts"] == []

    switched = await client.post(
        f"/inventory/voice/sessions/{session_id}/utterances",
        json={
            "client_event_id": "41111111-1111-4111-8111-111111111111",
            "sequence": 3,
            "transcript": "Switch to Voice Dry Storage",
        },
        headers=headers,
    )
    assert switched.status_code == 201, switched.text
    assert switched.json()["status"] == "ACCEPTED"

    dry_count = await client.post(
        f"/inventory/voice/sessions/{session_id}/utterances",
        json={
            "client_event_id": "51111111-1111-4111-8111-111111111111",
            "sequence": 4,
            "transcript": "Voice Tomatoes 5 pounds",
        },
        headers=headers,
    )
    assert dry_count.status_code == 201, dry_count.text

    finished = await client.post(
        f"/inventory/voice/sessions/{session_id}/finish", headers=headers
    )
    assert finished.status_code == 200, finished.text
    body = finished.json()
    assert body["status"] == "FINISHED"
    assert body["blocking_review_count"] == 0
    assert len(body["draft_counts"]) == 2
    effective = {
        (row["location_name"], row["item_name"]): row["quantity"]
        for row in body["effective_counts"]
    }
    assert effective[("Voice Walk In", "Voice Chicken Breast")] == "120.0000"
    assert effective[("Voice Dry Storage", "Voice Tomatoes")] == "5.0000"

    csv_export = await client.get(
        f"/inventory/voice/sessions/{session_id}/export.csv", headers=headers
    )
    assert csv_export.status_code == 200
    assert "Voice Chicken Breast" in csv_export.text
    assert "CORRECTED" in csv_export.text

    xlsx_export = await client.get(
        f"/inventory/voice/sessions/{session_id}/export.xlsx", headers=headers
    )
    assert xlsx_export.status_code == 200
    workbook = load_workbook(BytesIO(xlsx_export.content), read_only=True)
    assert workbook.sheetnames == ["Count Summary", "Voice Audit", "Needs Review"]
    summary_values = list(workbook["Count Summary"].values)
    assert any("Voice Chicken Breast" in row for row in summary_values)

    printable = await client.get(
        f"/inventory/voice/sessions/{session_id}/print", headers=headers
    )
    assert printable.status_code == 200
    assert "Voice Inventory Count" in printable.text

    submitted = await client.post(
        f"/inventory/counts/{body['draft_counts'][0]['inventory_count_id']}/submit",
        headers=headers,
    )
    assert submitted.status_code == 200, submitted.text


@pytest.mark.asyncio
async def test_unresolved_voice_entry_blocks_linked_count_submission(client):
    headers = await login_manager(client, "voice-blocking@example.com")
    location = await client.post(
        "/inventory/locations", json={"name": "Voice Blocking Location"}, headers=headers
    )
    item = await client.post(
        "/inventory/items",
        json={"name": "Voice Blocking Item", "base_unit": "each"},
        headers=headers,
    )
    session = await client.post(
        "/inventory/voice/sessions",
        json={
            "client_session_id": "61111111-1111-4111-8111-111111111111",
            "initial_location_id": location.json()["id"],
        },
        headers=headers,
    )
    session_id = session.json()["id"]
    accepted = await client.post(
        f"/inventory/voice/sessions/{session_id}/utterances",
        json={
            "client_event_id": "71111111-1111-4111-8111-111111111111",
            "sequence": 1,
            "transcript": "Voice Blocking Item 4 each",
        },
        headers=headers,
    )
    assert accepted.status_code == 201
    unresolved = await client.post(
        f"/inventory/voice/sessions/{session_id}/utterances",
        json={
            "client_event_id": "81111111-1111-4111-8111-111111111111",
            "sequence": 2,
            "transcript": "Mystery Product 2 boxes",
        },
        headers=headers,
    )
    assert unresolved.json()["status"] == "NEEDS_CLARIFICATION"
    finished = await client.post(
        f"/inventory/voice/sessions/{session_id}/finish", headers=headers
    )
    assert finished.json()["status"] == "NEEDS_REVIEW"
    count_id = finished.json()["draft_counts"][0]["inventory_count_id"]
    blocked = await client.post(
        f"/inventory/counts/{count_id}/submit", headers=headers
    )
    assert blocked.status_code == 409
    assert "unresolved" in blocked.json()["detail"]

    rejected = await client.patch(
        f"/inventory/voice/sessions/{session_id}/entries/"
        f"{unresolved.json()['entries'][0]['id']}",
        json={"review_status": "REJECTED"},
        headers=headers,
    )
    assert rejected.status_code == 200
    ready = await client.get(
        f"/inventory/voice/sessions/{session_id}", headers=headers
    )
    assert ready.json()["status"] == "FINISHED"
    submitted = await client.post(
        f"/inventory/counts/{count_id}/submit", headers=headers
    )
    assert submitted.status_code == 200, submitted.text
